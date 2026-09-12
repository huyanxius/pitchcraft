#!/usr/bin/env python3
"""Package 16:9 PNG slide images into an image-only PPTX."""

from __future__ import annotations

import argparse
import struct
import zipfile
from pathlib import Path


def png_size(path: Path) -> tuple[int, int]:
    with path.open("rb") as handle:
        header = handle.read(24)
    if header[:8] != b"\x89PNG\r\n\x1a\n":
        raise ValueError(f"Not a PNG: {path}")
    return struct.unpack(">II", header[16:24])


def build_contact_sheet(images: list[Path], output: Path) -> None:
    try:
        from PIL import Image, ImageDraw
    except ImportError:
        print("contact_sheet=skipped (Pillow unavailable)")
        return
    width, height, pad, label = 480, 270, 24, 28
    cols = 2
    rows = (len(images) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * (width + pad) + pad, rows * (height + label + pad) + pad), "white")
    draw = ImageDraw.Draw(sheet)
    for index, path in enumerate(images):
        image = Image.open(path).convert("RGB")
        image.thumbnail((width, height))
        x = pad + (index % cols) * (width + pad)
        y = pad + (index // cols) * (height + label + pad)
        draw.text((x, y), path.stem, fill=(30, 30, 30))
        sheet.paste(image, (x, y + label))
    output.parent.mkdir(parents=True, exist_ok=True)
    sheet.save(output)


def validate_pptx(path: Path) -> dict[str, int]:
    with zipfile.ZipFile(path) as archive:
        slides = sorted(name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml"))
        xml = "".join(archive.read(name).decode("utf-8", errors="ignore") for name in slides)
    return {
        "slides": len(slides),
        "pictures": xml.count("<p:pic>"),
        "editableTextBodies": xml.count("<p:txBody>"),
        "editableShapes": xml.count("<p:sp>"),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("image_dir", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--contact-sheet", type=Path)
    args = parser.parse_args()

    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise SystemExit("python-pptx is required. Use the bundled workspace Python runtime or install python-pptx.") from exc

    images = sorted(args.image_dir.glob("slide_*.png"))
    if not images:
        raise SystemExit(f"No slide_*.png files found in {args.image_dir}")
    for image in images:
        width, height = png_size(image)
        if abs(width / height - 16 / 9) > 0.01:
            raise SystemExit(f"{image.name} is not 16:9: {width}x{height}")

    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    while deck.slides:
        rel_id = deck.slides._sldIdLst[0].rId
        deck.part.drop_rel(rel_id)
        del deck.slides._sldIdLst[0]
    blank = deck.slide_layouts[6]
    for image in images:
        slide = deck.slides.add_slide(blank)
        slide.shapes.add_picture(str(image), 0, 0, width=deck.slide_width, height=deck.slide_height)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    deck.save(args.out)
    if args.contact_sheet:
        build_contact_sheet(images, args.contact_sheet)

    validation = validate_pptx(args.out)
    for key, value in validation.items():
        print(f"{key}={value}")
    if validation != {"slides": len(images), "pictures": len(images), "editableTextBodies": 0, "editableShapes": 0}:
        raise SystemExit("Image-only PPTX validation failed")


if __name__ == "__main__":
    main()
